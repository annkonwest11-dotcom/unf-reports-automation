#!/usr/bin/env python3
"""
Выгрузка взаиморасчётов в Excel
"""

import requests
import base64
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from collections import defaultdict

LOGIN = "api_bot"
PASSWORD = "slavaperfilev1414"
BASE_ID = "152757"
BASE_URL = f"https://base.42clouds.com/unf/{BASE_ID}/odata/standard.odata/"

START_DATE = "2026-06-17"
END_DATE = "2026-06-24"

def get_basic_auth():
    creds = base64.b64encode(f"{LOGIN}:{PASSWORD}".encode()).decode()
    return {"Authorization": f"Basic {creds}", "Accept": "application/json"}

def fetch_odata(entity):
    url = f"{BASE_URL}{entity}"
    headers = get_basic_auth()
    results = []
    while url:
        resp = requests.get(url, headers=headers, timeout=30, verify=False)
        if resp.status_code != 200:
            return None
        data = resp.json()
        if 'value' in data:
            results.extend(data['value'])
        url = data.get('@odata.nextLink')
    return results

def get_contractor_names():
    contractors = fetch_odata('Catalog_Контрагенты')
    if not contractors:
        return {}
    names_map = {}
    for c in contractors:
        ref_key = c.get('Ref_Key')
        description = c.get('Description', 'Unknown')
        names_map[ref_key] = description
    return names_map

import urllib3
urllib3.disable_warnings()

print("Загружаю взаиморасчёты...")

contractor_names = get_contractor_names()
settlements = fetch_odata('AccumulationRegister_РасчетыСПокупателями_RecordType')

if not settlements:
    print("❌ Нет данных")
    exit(1)

# Считаем начальный остаток (до START_DATE)
start_date_obj = datetime.strptime(START_DATE, '%Y-%m-%d')
end_date_obj = datetime.strptime(END_DATE, '%Y-%m-%d')

initial_balance = defaultdict(float)  # До START_DATE
changes_increase = defaultdict(float)  # Увеличение (Receipt)
changes_decrease = defaultdict(float)  # Уменьшение (Expense)

for s in settlements:
    period_str = s.get('Period', '')
    date_str = period_str[:10]

    ctg_key = s.get('Контрагент_Key', 'unknown')
    amount = s.get('Сумма', 0) or 0
    record_type = s.get('RecordType', 'Receipt')

    # Считаем начальный остаток (всё до START_DATE)
    if date_str < START_DATE:
        if record_type == 'Receipt':
            initial_balance[ctg_key] += amount
        elif record_type == 'Expense':
            initial_balance[ctg_key] -= amount

    # Считаем изменения в периоде
    elif START_DATE <= date_str <= END_DATE:
        if record_type == 'Receipt':
            changes_increase[ctg_key] += amount
        elif record_type == 'Expense':
            changes_decrease[ctg_key] += amount

# Создаём отчёт
report_data = []

for ctg_key in sorted(set(list(initial_balance.keys()) + list(changes_increase.keys()) + list(changes_decrease.keys()))):
    name = contractor_names.get(ctg_key, 'Unknown')
    initial = initial_balance.get(ctg_key, 0)
    increase = changes_increase.get(ctg_key, 0)
    decrease = changes_decrease.get(ctg_key, 0)
    final = initial + increase - decrease

    # Выводим только если есть движения
    if initial != 0 or increase != 0 or decrease != 0 or final != 0:
        report_data.append({
            'Контрагент': name,
            'Начальный остаток': initial,
            'Увеличение долга': increase,
            'Уменьшение долга': decrease,
            'Конечный остаток': final
        })

# Создаём Excel
wb = openpyxl.Workbook()
ws = wb.active
ws.title = f"Взаиморасчёты {START_DATE}"

# Заголовки
headers = ['Контрагент', 'Начальный остаток', 'Увеличение долга', 'Уменьшение долга', 'Конечный остаток']
ws.append(headers)

# Стиль заголовка
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF")
for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

# Добавляем данные
for row in report_data:
    ws.append([
        row['Контрагент'],
        row['Начальный остаток'],
        row['Увеличение долга'],
        row['Уменьшение долга'],
        row['Конечный остаток']
    ])

# Форматирование
thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=5):
    for cell in row:
        cell.border = thin_border
        if cell.row > 1:  # Не заголовок
            if cell.column > 1:  # Числовые колонки
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal="right")
            else:  # Контрагент
                cell.alignment = Alignment(horizontal="left")

# Ширины колонок
ws.column_dimensions['A'].width = 50
ws.column_dimensions['B'].width = 18
ws.column_dimensions['C'].width = 18
ws.column_dimensions['D'].width = 18
ws.column_dimensions['E'].width = 18

# ИТОГО строка
total_row = ws.max_row + 2
ws[f'A{total_row}'] = 'ИТОГО'
ws[f'A{total_row}'].font = Font(bold=True)

ws[f'B{total_row}'] = f'=SUM(B2:B{total_row-2})'
ws[f'C{total_row}'] = f'=SUM(C2:C{total_row-2})'
ws[f'D{total_row}'] = f'=SUM(D2:D{total_row-2})'
ws[f'E{total_row}'] = f'=SUM(E2:E{total_row-2})'

for cell in ws[total_row]:
    cell.font = Font(bold=True)
    if cell.column > 1:
        cell.number_format = '#,##0.00'

file_path = f"/Users/anna/claude-test/Взаиморасчёты_{START_DATE}_{END_DATE}.xlsx"
wb.save(file_path)

print(f"\n✅ Файл создан: {file_path}")
print(f"✅ Контрагентов в отчёте: {len(report_data)}")

# Вычисляем ИТОГО
total_initial = sum(r['Начальный остаток'] for r in report_data)
total_increase = sum(r['Увеличение долга'] for r in report_data)
total_decrease = sum(r['Уменьшение долга'] for r in report_data)
total_final = sum(r['Конечный остаток'] for r in report_data)

print(f"\n📊 ИТОГО:")
print(f"  Начальный остаток: {total_initial:,.2f} ₽")
print(f"  Увеличение долга: {total_increase:,.2f} ₽")
print(f"  Уменьшение долга: {total_decrease:,.2f} ₽")
print(f"  Конечный остаток: {total_final:,.2f} ₽")
