# -*- coding: utf-8 -*-
"""Отчёт по дебиторке: все ИП, все должники, топ просроченных → PDF + XLSX на Desktop.

Запуск:  python3 debt_report.py            (срез на сегодня)
         python3 debt_report.py 2026-10-01 (срез на дату)

Чем отличается от sync_debitorka.py: тот распределяет долги ПО МЕНЕДЖЕРАМ в
Google-таблицу и ставит задачи в Битрикс, а здесь — разрез ПО ИП (три базы 1С)
с полными списками должников и общим топом просрочки, файлами для печати.
Методика расчёта та же (fetch_base из sync_debitorka): долг = расходные накладные
с положительным остатком, просрочка — только при заданной отсрочке договора.
"""
import datetime
import json
import os
import subprocess
import sys
from collections import defaultdict

sys.path.insert(0, "/Users/anna/GREENCH/01-деньги/зарплата")
os.chdir("/Users/anna/GREENCH/01-деньги/зарплата")

from sync_debitorka import fetch_base            # noqa: E402
from sync_odata import BASES                      # noqa: E402
from debt_payments import by_name, бренд           # noqa: E402

import tempfile

SCRATCH = tempfile.gettempdir()      # временный html для печати
DESKTOP = os.path.expanduser("~/Desktop")

ИП = {                       # база 1С → как называть ИП в отчёте
    "perfilev": "ИП Перфильев В. К.",
    "gubarev": "ИП Губарев И. С.",
    "volodihina": "ИП Володихина В. В.",
}
КОРЗИНЫ = ["До 7 дн.", "8–14 дн.", "15–30 дн.", "31–60 дн.", "От 61 дн."]

ОКНО_ОПЛАТ = 30          # дней для колонки «оплачено за период»
ОКНО_ПОСЛЕДНИХ = 14     # дней для листа/страницы «последние оплаты»

СРЕЗ = (datetime.date.fromisoformat(sys.argv[1]) if len(sys.argv) > 1
        else datetime.date.today())

итог = {}
for ключ, cfg in BASES.items():
    try:
        записи = fetch_base(cfg["id"], СРЕЗ)
    except Exception as e:                                    # noqa: BLE001
        print("база %s (%s) — ОШИБКА: %s" % (ключ, cfg["id"], str(e)[:200]))
        continue
    try:
        оплаты, висят = by_name(cfg["id"], СРЕЗ)
    except Exception as e:                                    # noqa: BLE001
        print("база %s — оплаты НЕ ВЫГРУЖЕНЫ: %s" % (ключ, str(e)[:200]))
        raise                        # без оплат отчёт собирать нельзя — это пустые колонки
    итог[ключ] = {"id": cfg["id"], "records": записи,
                  "payments": оплаты, "unapplied": висят}
    долг = sum(r["total"] for r in записи)
    просроч = sum(r["overdue"] for r in записи)
    print("база %-12s (%s): должников %3d | долг %12.2f | просрочка %12.2f"
          " | оплат %d | без разноски %12.2f"
          % (ключ, cfg["id"], len(записи), долг, просроч,
             sum(len(v) for v in оплаты.values()), sum(sum(v) for v in висят.values())))

сырое = {"asof": СРЕЗ.isoformat(), "bases": итог}

# Долг посчитан на СРЕЗ 00:00 (так же, как в отчёте 1С «по срокам долга»), поэтому
# оплата, пришедшая в день среза, в долге ещё НЕ учтена — помечаем такие отдельно.
МОМЕНТ = datetime.datetime.combine(СРЕЗ, datetime.time.min)
КОНЕЦ_ДНЯ = datetime.datetime.combine(СРЕЗ, datetime.time.max)
НАЧАЛО_ОКНА = МОМЕНТ - datetime.timedelta(days=ОКНО_ОПЛАТ)
НАЧАЛО_ПОСЛЕДНИХ = МОМЕНТ - datetime.timedelta(days=ОКНО_ПОСЛЕДНИХ)

# Клиента переводят с юрлица на юрлицо и заводят ему вторую карточку: накладные
# идут на новую, а деньги продолжают приходить на старую. По своей карточке у него
# «нет оплат», хотя платит он исправно — поэтому ищем ещё и по бренду.
оплаты_бренда = defaultdict(list)
for ключ, пакет in сырое["bases"].items():
    for имя, список in пакет["payments"].items():
        for дата, сумма, куда in список:
            оплаты_бренда[бренд(имя)].append(
                (дата, сумма, куда, имя, ИП.get(ключ, ключ)))
for v in оплаты_бренда.values():
    v.sort(key=lambda x: x[0])

строки = []
for ключ, пакет in сырое["bases"].items():
    for r in пакет["records"]:
        история = [о for о in пакет["payments"].get(r["name"], []) if о[0] <= КОНЕЦ_ДНЯ]
        карточка = ""
        if not история:
            чужие = [о for о in оплаты_бренда.get(бренд(r["name"]), [])
                     if о[3] != r["name"] and о[0] <= КОНЕЦ_ДНЯ]
            if чужие:
                история = [(д, с, к) for д, с, к, _, _ in чужие]
                карточка = f"{чужие[-1][3]} ({чужие[-1][4].replace('ИП ', '')})"
        посл = история[-1] if история else None
        строки.append({
            "ип": ИП.get(ключ, ключ), "name": r["name"], "total": r["total"],
            "overdue": r["overdue"], "days": r["days"], "buckets": r["buckets"],
            "pay_date": посл[0] if посл else None,
            "pay_sum": посл[1] if посл else 0.0,
            "pay_where": посл[2] if посл else "",
            "pay_30": round(sum(о[1] for о in история if о[0] >= НАЧАЛО_ОКНА), 2),
            "pay_cnt30": sum(1 for о in история if о[0] >= НАЧАЛО_ОКНА),
            "no_pay_days": ((СРЕЗ - посл[0].date()).days if посл else None),
            "pay_card": карточка,
            "unapplied": sum(пакет["unapplied"].get(r["name"], (0.0, 0.0))),
            "un_free": пакет["unapplied"].get(r["name"], (0.0, 0.0))[0],
            "un_over": пакет["unapplied"].get(r["name"], (0.0, 0.0))[1],
        })

# все оплаты за последние ОКНО_ПОСЛЕДНИХ дней — отдельным списком
долг_по_клиенту = {(s["ип"], s["name"]): s for s in строки}
оплаты_список = []
for ключ, пакет in сырое["bases"].items():
    ип = ИП.get(ключ, ключ)
    for имя, список in пакет["payments"].items():
        for дата, сумма, куда in список:
            if дата < НАЧАЛО_ПОСЛЕДНИХ or дата > КОНЕЦ_ДНЯ:
                continue
            долж = долг_по_клиенту.get((ип, имя))
            оплаты_список.append({
                "дата": дата, "name": имя, "ип": ип, "сумма": сумма, "куда": куда,
                "в_долге": дата < МОМЕНТ,
                "долг": долж["total"] if долж else 0.0,
            })
оплаты_список.sort(key=lambda x: (-x["дата"].timestamp(), -x["сумма"]))
оплачено_окно = sum(s["pay_30"] for s in строки)

# деньги без разноски — в т.ч. у тех, кто в списке должников
без_разноски = []
for ключ, пакет in сырое["bases"].items():
    ип = ИП.get(ключ, ключ)
    for имя, (оплата, переплата) in пакет["unapplied"].items():
        долж = долг_по_клиенту.get((ип, имя))
        без_разноски.append({
            "name": имя, "ип": ип, "сумма": оплата + переплата,
            "оплата": оплата, "переплата": переплата,
            "долг": долж["total"] if долж else 0.0,
            "просроч": долж["overdue"] if долж else 0.0,
        })
без_разноски.sort(key=lambda x: -min(x["сумма"], x["долг"]) if x["долг"] else 1)
висит_всего = sum(x["сумма"] for x in без_разноски)
снимет_разноска = sum(min(x["сумма"], x["долг"]) for x in без_разноски)
должников_с_деньгами = sum(1 for x in без_разноски if x["долг"] > 0)

по_ип = defaultdict(list)
for s in строки:
    по_ип[s["ип"]].append(s)
for v in по_ип.values():
    v.sort(key=lambda r: -r["total"])

просроченные = sorted([s for s in строки if s["overdue"] > 0],
                      key=lambda r: -r["overdue"])
долг_всего = sum(s["total"] for s in строки)
просроч_всего = sum(s["overdue"] for s in строки)

# клиенты, которые должны сразу нескольким ИП
счёт = defaultdict(set)
for s in строки:
    счёт[s["name"]].add(s["ип"])
двойные = sorted([(n, sorted(ипы)) for n, ипы in счёт.items() if len(ипы) > 1])


def руб(x):
    return f"{x:,.0f}".replace(",", " ")


def дата_кратко(d):
    return f"{d:%d.%m}" if d else "—"


# оплаты и «висящие» деньги в разрезе ИП — по ВСЕМ клиентам базы, не только должникам
оплаты_по_ип_окно = defaultdict(float)
оплаты_по_ип_кол = defaultdict(int)
висит_по_ип = defaultdict(float)
for _ключ, _пакет in сырое["bases"].items():
    _ип = ИП.get(_ключ, _ключ)
    висит_по_ип[_ип] += sum(sum(_v) for _v in _пакет["unapplied"].values())
    for _список in _пакет["payments"].values():
        for _дата, _сумма, _ in _список:
            if НАЧАЛО_ОКНА <= _дата <= КОНЕЦ_ДНЯ:
                оплаты_по_ип_окно[_ип] += _сумма
                оплаты_по_ип_кол[_ип] += 1
оплачено_все_окно = sum(оплаты_по_ип_окно.values())
оплат_всего_окно = sum(оплаты_по_ип_кол.values())


# ────────────────────────────── XLSX ──────────────────────────────
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ЗЕЛ = "2E7D32"
ШАПКА_ФОН = "EEF5EE"
КРАСН = "FFEBEE"

wb = Workbook()


def оформить_шапку(ws, заголовки, строка=1):
    ws.append(заголовки)
    for c in range(1, len(заголовки) + 1):
        я = ws.cell(row=строка, column=c)
        я.font = Font(bold=True, color="2E5B31")
        я.fill = PatternFill("solid", fgColor=ШАПКА_ФОН)
        я.border = Border(bottom=Side("medium", color="CFE0CF"))
        я.alignment = Alignment(wrap_text=True, vertical="center")
    # строкой, а не ws.cell(): обращение к ячейке создало бы её и сдвинуло
    # первую строку данных вниз — между шапкой и данными висела пустая
    ws.freeze_panes = f"A{строка + 1}"


# лист СВОДКА
ws = wb.active
ws.title = "СВОДКА"
ws["A1"] = "Дебиторская задолженность по всем ИП"
ws["A1"].font = Font(size=15, bold=True, color=ЗЕЛ)
ws["A2"] = f"Срез на {СРЕЗ:%d.%m.%Y} · источник: 1С (OData), три базы"
ws["A2"].font = Font(size=10, color="6B7280")
ws.append([])
ws.append([])
оформить_шапку(ws, ["ИП", "Должников", "Общий долг, ₽", "В т.ч. просрочено, ₽",
                    "Доля просрочки", "Должников с просрочкой",
                    f"Оплат за {ОКНО_ОПЛАТ} дн., шт.", f"Оплачено за {ОКНО_ОПЛАТ} дн., ₽",
                    "Деньги без разноски, ₽"], строка=4)
for ип in sorted(по_ип, key=lambda k: -sum(r["total"] for r in по_ип[k])):
    записи = по_ип[ип]
    долг = sum(r["total"] for r in записи)
    про = sum(r["overdue"] for r in записи)
    ws.append([ип, len(записи), долг, про, (про / долг if долг else 0),
               sum(1 for r in записи if r["overdue"] > 0),
               оплаты_по_ип_кол[ип], оплаты_по_ип_окно[ип], висит_по_ип[ип]])
итог = ws.max_row + 1
ws.append(["ИТОГО", len(строки), долг_всего, просроч_всего,
           (просроч_всего / долг_всего if долг_всего else 0), len(просроченные),
           оплат_всего_окно, оплачено_все_окно, висит_всего])
for c in range(1, 10):
    ws.cell(row=итог, column=c).font = Font(bold=True)
    ws.cell(row=итог, column=c).fill = PatternFill("solid", fgColor=ШАПКА_ФОН)
for строка_ in ws.iter_rows(min_row=5, max_row=ws.max_row):
    строка_[2].number_format = '# ##0'
    строка_[3].number_format = '# ##0'
    строка_[4].number_format = '0.0%'
    строка_[7].number_format = '# ##0'
    строка_[8].number_format = '# ##0'
ряд = ws.max_row + 2
ws.cell(row=ряд, column=1, value="Оплаты — проведённые поступления от покупателей "
        f"(счёт и касса) за {ОКНО_ОПЛАТ} дн. до среза, по всем клиентам базы.")
ws.cell(row=ряд, column=1).font = Font(size=9, color="6B7280")
ws.cell(row=ряд + 1, column=1, value="«Деньги без разноски» — оплаты и переплаты, "
        "не разнесённые по накладным: долг по ним висит, хотя клиент заплатил. "
        f"Разноска снимет {руб(снимет_разноска)} ₽ у {должников_с_деньгами} должников.")
ws.cell(row=ряд + 1, column=1).font = Font(size=9, color="6B7280")
for кол, ш in zip("ABCDEFGHI", (26, 11, 16, 20, 14, 22, 18, 20, 22)):
    ws.column_dimensions[кол].width = ш

# лист ТОП ПРОСРОЧКИ
ws = wb.create_sheet("ТОП ПРОСРОЧКИ")
ws["A1"] = "Должники с просроченной задолженностью — по убыванию суммы"
ws["A1"].font = Font(size=13, bold=True, color=ЗЕЛ)
ws.append([])
оформить_шапку(ws, ["№", "Контрагент", "ИП", "Просрочено, ₽", "Общий долг, ₽",
                    "Дней просрочки"] + КОРЗИНЫ
                   + ["Посл. оплата", "Сумма посл. оплаты, ₽", "Дней без оплаты",
                      f"Оплачено за {ОКНО_ОПЛАТ} дн., ₽", "Деньги без разноски, ₽",
                      "Оплата на чужую карточку"],
               строка=3)
for n, s in enumerate(просроченные, 1):
    ws.append([n, s["name"], s["ип"], s["overdue"], s["total"], s["days"]] + s["buckets"]
              + [дата_кратко(s["pay_date"]), s["pay_sum"] or None,
                 s["no_pay_days"] if s["no_pay_days"] is not None else "нет оплат",
                 s["pay_30"] or None, s["unapplied"] or None, s["pay_card"] or None])
    if n <= 10:
        for c in range(1, 18):
            ws.cell(row=ws.max_row, column=c).fill = PatternFill("solid", fgColor=КРАСН)
    if s["unapplied"] > 0:
        ws.cell(row=ws.max_row, column=16).font = Font(color="B26A00", bold=True)
for строка_ in ws.iter_rows(min_row=4, max_row=ws.max_row):
    for i in (3, 4):
        строка_[i].number_format = '# ##0'
    for i in range(6, 11):
        строка_[i].number_format = '# ##0'
    for i in (12, 14, 15):
        строка_[i].number_format = '# ##0'
    строка_[5].alignment = Alignment(horizontal="center")
    строка_[11].alignment = Alignment(horizontal="center")
    строка_[13].alignment = Alignment(horizontal="center")
for кол, ш in zip("ABCDEFGHIJKLMNOPQ",
                  (5, 46, 22, 15, 15, 14, 12, 12, 12, 12, 12, 13, 20, 16, 20, 22, 46)):
    ws.column_dimensions[кол].width = ш

# по листу на ИП
for ип in sorted(по_ип, key=lambda k: -sum(r["total"] for r in по_ип[k])):
    записи = по_ип[ип]
    ws = wb.create_sheet(ип.replace("ИП ", "")[:28])
    ws["A1"] = f"{ип} — все должники на {СРЕЗ:%d.%m.%Y}"
    ws["A1"].font = Font(size=13, bold=True, color=ЗЕЛ)
    ws.append([])
    оформить_шапку(ws, ["№", "Контрагент", "Общий долг, ₽", "Просрочено, ₽",
                        "Дней просрочки"] + КОРЗИНЫ
                       + ["Посл. оплата", "Сумма посл. оплаты, ₽", "Дней без оплаты",
                          f"Оплачено за {ОКНО_ОПЛАТ} дн., ₽", "Деньги без разноски, ₽",
                          "Оплата на чужую карточку"],
                   строка=3)
    for n, s in enumerate(записи, 1):
        ws.append([n, s["name"], s["total"], s["overdue"], s["days"]] + s["buckets"]
                  + [дата_кратко(s["pay_date"]), s["pay_sum"] or None,
                     s["no_pay_days"] if s["no_pay_days"] is not None else "нет оплат",
                     s["pay_30"] or None, s["unapplied"] or None, s["pay_card"] or None])
        if s["overdue"] > 0:
            ws.cell(row=ws.max_row, column=4).font = Font(color="C62828", bold=True)
        if s["unapplied"] > 0:
            ws.cell(row=ws.max_row, column=15).font = Font(color="B26A00", bold=True)
    строка_итога = ws.max_row + 1
    ws.append(["", "ИТОГО", f"=SUM(C4:C{ws.max_row})", f"=SUM(D4:D{ws.max_row})", ""]
              + [f"=SUM({get_column_letter(6 + i)}4:{get_column_letter(6 + i)}{ws.max_row})"
                 for i in range(5)]
              + ["", "", "", f"=SUM(N4:N{ws.max_row})", f"=SUM(O4:O{ws.max_row})"])
    for c in range(1, 17):
        ws.cell(row=строка_итога, column=c).font = Font(bold=True)
        ws.cell(row=строка_итога, column=c).fill = PatternFill("solid", fgColor=ШАПКА_ФОН)
    for строка_ in ws.iter_rows(min_row=4, max_row=ws.max_row):
        for i in (2, 3):
            строка_[i].number_format = '# ##0'
        for i in range(5, 10):
            строка_[i].number_format = '# ##0'
        for i in (11, 13, 14):
            строка_[i].number_format = '# ##0'
        строка_[4].alignment = Alignment(horizontal="center")
        строка_[10].alignment = Alignment(horizontal="center")
        строка_[12].alignment = Alignment(horizontal="center")
    for кол, ш in zip("ABCDEFGHIJKLMNOP",
                      (5, 48, 15, 15, 14, 12, 12, 12, 12, 12, 13, 20, 16, 20, 22, 46)):
        ws.column_dimensions[кол].width = ш


# лист ПОСЛЕДНИЕ ОПЛАТЫ
ws = wb.create_sheet("ПОСЛЕДНИЕ ОПЛАТЫ")
ws["A1"] = f"Оплаты от покупателей за последние {ОКНО_ПОСЛЕДНИХ} дн. (по {СРЕЗ:%d.%m.%Y})"
ws["A1"].font = Font(size=13, bold=True, color=ЗЕЛ)
ws["A2"] = ("проведённые поступления на счёт и в кассу; «учтено в долге» = оплата прошла "
            "до момента среза, иначе долг выше фактического")
ws["A2"].font = Font(size=10, color="6B7280")
оформить_шапку(ws, ["Дата", "Контрагент", "ИП", "Сумма, ₽", "Куда",
                    "Учтено в долге", "Долг сейчас, ₽"], строка=3)
for о in оплаты_список:
    ws.append([о["дата"].strftime("%d.%m.%Y %H:%M"), о["name"], о["ип"].replace("ИП ", ""),
               о["сумма"], о["куда"], "да" if о["в_долге"] else "нет — после среза",
               о["долг"] or None])
    if not о["в_долге"]:
        for c in range(1, 8):
            ws.cell(row=ws.max_row, column=c).fill = PatternFill("solid", fgColor="FFF8E1")
строка_итога = ws.max_row + 1
ws.append(["", "ИТОГО", f"{len(оплаты_список)} оплат",
           f"=SUM(D4:D{ws.max_row})", "", "", ""])
for c in range(1, 8):
    ws.cell(row=строка_итога, column=c).font = Font(bold=True)
    ws.cell(row=строка_итога, column=c).fill = PatternFill("solid", fgColor=ШАПКА_ФОН)
for строка_ in ws.iter_rows(min_row=4, max_row=ws.max_row):
    строка_[3].number_format = '# ##0'
    строка_[6].number_format = '# ##0'
for кол, ш in zip("ABCDEFG", (18, 48, 16, 15, 10, 20, 16)):
    ws.column_dimensions[кол].width = ш

# лист ДЕНЬГИ БЕЗ РАЗНОСКИ
ws = wb.create_sheet("ДЕНЬГИ БЕЗ РАЗНОСКИ")
ws["A1"] = f"Оплаты и переплаты, не разнесённые по накладным — на {СРЕЗ:%d.%m.%Y}"
ws["A1"].font = Font(size=13, bold=True, color=ЗЕЛ)
ws["A2"] = (f"висит {руб(висит_всего)} ₽ у {len(без_разноски)} контрагентов; "
            f"у {должников_с_деньгами} из них есть долг — разноска снимет "
            f"{руб(снимет_разноска)} ₽. До разноски такого клиента НЕ звонить.")
ws["A2"].font = Font(size=10, color="6B7280")
оформить_шапку(ws, ["Контрагент", "ИП", "Денег без разноски, ₽",
                    "В т.ч. неразнесённая оплата, ₽", "В т.ч. переплата по накладной, ₽",
                    "Долг, ₽", "В т.ч. просрочено, ₽", "Снимет разноска, ₽",
                    "Что делать"], строка=3)
for x in без_разноски:
    снимет = min(x["сумма"], x["долг"])
    если = ("разнести — долг закроется" if x["долг"] and x["сумма"] >= x["долг"]
            else "разнести — долг уменьшится" if x["долг"]
            else "аванс/переплата, долга нет")
    ws.append([x["name"], x["ип"].replace("ИП ", ""), x["сумма"],
               x["оплата"] or None, x["переплата"] or None,
               x["долг"] or None, x["просроч"] or None, снимет or None, если])
    if x["долг"] and x["сумма"] >= x["долг"]:
        for c in range(1, 10):
            ws.cell(row=ws.max_row, column=c).fill = PatternFill("solid", fgColor="FFF8E1")
строка_итога = ws.max_row + 1
ws.append(["ИТОГО", ""] + [f"=SUM({б}4:{б}{ws.max_row})" for б in "CDEFGH"] + [""])
for c in range(1, 10):
    ws.cell(row=строка_итога, column=c).font = Font(bold=True)
    ws.cell(row=строка_итога, column=c).fill = PatternFill("solid", fgColor=ШАПКА_ФОН)
for строка_ in ws.iter_rows(min_row=4, max_row=ws.max_row):
    for i in range(2, 8):
        строка_[i].number_format = '# ##0'
for кол, ш in zip("ABCDEFGHI", (48, 16, 22, 26, 28, 15, 20, 20, 30)):
    ws.column_dimensions[кол].width = ш

ИМЯ = f"Дебиторка по ИП {СРЕЗ:%d.%m.%Y}"
xlsx_path = f"{DESKTOP}/{ИМЯ}.xlsx"
wb.save(xlsx_path)
print("XLSX:", xlsx_path)


# ────────────────────────────── PDF ──────────────────────────────
def оплата_ячейка(s):
    """«12.10 · 48 509» — когда и сколько клиент заплатил в последний раз."""
    if not s["pay_date"]:
        return '<td class="pay">нет оплат</td>'
    return (f'<td class="pay">{дата_кратко(s["pay_date"])} · '
            f'{руб(s["pay_sum"])}</td>')


def метка_денег(s):
    """Клиент заплатил, но деньги не разнесены — звонить ему нельзя."""
    куски = []
    if s.get("unapplied"):
        куски.append(f'<div class="flag">⚑ не разнесено {руб(s["unapplied"])} ₽</div>')
    if s.get("pay_card"):
        куски.append(f'<div class="card2">платит на карточку «{s["pay_card"]}»</div>')
    return "".join(куски)


def таблица_топ(записи, лимит=None):
    куски = []
    for n, s in enumerate(записи[:лимит] if лимит else записи, 1):
        класс = ' class="hot"' if n <= 10 else ""
        куски.append(
            f'<tr{класс}><td class="n">{n}</td><td>{s["name"]}{метка_денег(s)}</td>'
            f'<td class="ip">{s["ип"].replace("ИП ", "")}</td>'
            f'<td class="p">{руб(s["overdue"])}</td><td class="p">{руб(s["total"])}</td>'
            f'<td class="d">{s["days"]}</td>{оплата_ячейка(s)}</tr>')
    return "\n".join(куски)


def таблица_ип(записи):
    куски = []
    for n, s in enumerate(записи, 1):
        про = f'<span class="red">{руб(s["overdue"])}</span>' if s["overdue"] else "—"
        куски.append(
            f'<tr><td class="n">{n}</td><td>{s["name"]}{метка_денег(s)}</td>'
            f'<td class="p">{руб(s["total"])}</td><td class="p">{про}</td>'
            f'<td class="d">{s["days"] or "—"}</td>{оплата_ячейка(s)}</tr>')
    return "\n".join(куски)


def таблица_оплат(записи, лимит=None):
    куски = []
    for о in (записи[:лимит] if лимит else записи):
        класс = "" if о["в_долге"] else ' class="late"'
        куски.append(
            f'<tr{класс}><td class="d">{о["дата"]:%d.%m}</td><td>{о["name"]}</td>'
            f'<td class="ip">{о["ип"].replace("ИП ", "")}</td>'
            f'<td class="p">{руб(о["сумма"])}</td>'
            f'<td class="d">{о["куда"]}</td>'
            f'<td class="p">{руб(о["долг"]) if о["долг"] else "—"}</td></tr>')
    return "\n".join(куски)


def таблица_без_разноски(записи):
    куски = []
    for x in записи:
        снимет = min(x["сумма"], x["долг"])
        класс = ' class="hot"' if x["долг"] and x["сумма"] >= x["долг"] else ""
        куски.append(
            f'<tr{класс}><td>{x["name"]}</td>'
            f'<td class="ip">{x["ип"].replace("ИП ", "")}</td>'
            f'<td class="p">{руб(x["сумма"])}</td>'
            f'<td class="p">{руб(x["переплата"]) if x["переплата"] else "—"}</td>'
            f'<td class="p">{руб(x["долг"]) if x["долг"] else "—"}</td>'
            f'<td class="p">{руб(снимет) if снимет else "—"}</td></tr>')
    return "\n".join(куски)


сводка_строки = ""
for ип in sorted(по_ип, key=lambda k: -sum(r["total"] for r in по_ип[k])):
    записи = по_ип[ип]
    долг = sum(r["total"] for r in записи)
    про = sum(r["overdue"] for r in записи)
    сводка_строки += (
        f'<tr><td>{ип}</td><td class="d">{len(записи)}</td>'
        f'<td class="p">{руб(долг)}</td><td class="p red">{руб(про)}</td>'
        f'<td class="d">{про / долг * 100:.0f}%</td>'
        f'<td class="p">{руб(оплаты_по_ип_окно[ип])}</td>'
        f'<td class="p">{руб(висит_по_ип[ип]) if висит_по_ип[ип] else "—"}</td></tr>')

страницы_ип = ""
for ип in sorted(по_ип, key=lambda k: -sum(r["total"] for r in по_ип[k])):
    записи = по_ип[ип]
    страницы_ип += f"""
  <div class="page">
    <h2>{ип} — все должники ({len(записи)})</h2>
    <table><thead><tr><th class="n">№</th><th>Контрагент</th><th class="p">Долг, ₽</th>
      <th class="p">Просрочено, ₽</th><th class="d">Дней</th>
      <th class="pay">Посл. оплата</th></tr></thead>
      <tbody>{таблица_ип(записи)}</tbody></table>
  </div>"""

двойные_блок = ""
if двойные:
    пункты = "".join(f"<li>{n} — {', '.join(и.replace('ИП ', '') for и in ипы)}</li>"
                     for n, ипы in двойные[:20])
    двойные_блок = (f'<h3>Должны сразу нескольким ИП ({len(двойные)})</h3>'
                    f'<ul class="multi">{пункты}</ul>')

HTML = f"""<!doctype html>
<html lang="ru"><head><meta charset="utf-8"><title>{ИМЯ}</title>
<style>
  @page {{ size: A4; margin: 10mm 11mm; }}
  body {{ font-family: -apple-system, "Helvetica Neue", Arial, sans-serif;
         color: #1f2a37; font-size: 9pt; }}
  .head {{ display: flex; justify-content: space-between; align-items: flex-end;
          border-bottom: 3px solid #2e7d32; padding-bottom: 6px; margin-bottom: 10px; }}
  .brand {{ font-size: 17pt; font-weight: 700; color: #2e7d32; }}
  .sub {{ color: #6b7280; font-size: 8.5pt; margin-top: 2px; }}
  .right {{ text-align: right; }}
  .right b {{ font-size: 12pt; }}
  h2 {{ font-size: 12pt; color: #2e5b31; margin: 0 0 7px; }}
  h3 {{ font-size: 10pt; color: #2e5b31; margin: 12px 0 5px; }}
  table {{ width: 100%; border-collapse: collapse; margin-bottom: 4px; }}
  thead th {{ background: #eef5ee; color: #2e5b31; font-size: 8pt; text-transform: uppercase;
             letter-spacing: .3px; text-align: left; padding: 4px 7px;
             border-bottom: 2px solid #cfe0cf; }}
  td {{ padding: 2.6px 7px; border-bottom: 1px solid #eef1f3; }}
  tbody tr:nth-child(even) td {{ background: #fafbfa; }}
  .n {{ width: 26px; color: #9aa3ad; }}
  .p {{ width: 95px; text-align: right; font-variant-numeric: tabular-nums; font-weight: 600; }}
  thead th.p {{ text-align: right; }}
  .d {{ width: 56px; text-align: center; color: #6b7280; }}
  thead th.d {{ text-align: center; }}
  .ip {{ width: 92px; color: #6b7280; }}
  .pay {{ width: 112px; text-align: right; color: #374151;
         font-variant-numeric: tabular-nums; }}
  thead th.pay {{ text-align: right; }}
  .flag {{ color: #b26a00; font-size: 7.3pt; margin-top: 1px; }}
  .card2 {{ color: #46617a; font-size: 7.3pt; margin-top: 1px; }}
  .late td {{ background: #fff8e1 !important; }}
  .note {{ color: #6b7280; font-size: 8.5pt; margin: 2px 0 8px; }}
  .red {{ color: #c62828; }}
  .hot td {{ background: #ffebee !important; }}
  .itogo td {{ background: #eef5ee !important; font-weight: 700; }}
  .cards {{ display: flex; gap: 8px; margin: 10px 0 14px; }}
  .card {{ flex: 1; border: 1px solid #e3e8e3; border-radius: 6px; padding: 8px 10px; }}
  .card .v {{ font-size: 15pt; font-weight: 700; color: #2e7d32; }}
  .card .v.red {{ color: #c62828; }}
  .card .v.amber {{ color: #b26a00; }}
  .card .k {{ color: #6b7280; font-size: 8pt; }}
  .multi {{ columns: 2; font-size: 8.5pt; color: #374151; margin: 0; padding-left: 16px; }}
  .page {{ page-break-before: always; }}
</style></head>
<body>
  <div class="head">
    <div><div class="brand">GREENCH</div>
      <div class="sub">дебиторская задолженность · все ИП</div></div>
    <div class="right"><div>Отчёт на</div><b>{СРЕЗ:%d.%m.%Y}</b>
      <div class="sub">источник: 1С, три базы</div></div>
  </div>

  <div class="cards">
    <div class="card"><div class="k">Должников всего</div><div class="v">{len(строки)}</div></div>
    <div class="card"><div class="k">Общий долг</div><div class="v">{руб(долг_всего)} ₽</div></div>
    <div class="card"><div class="k">Просрочено</div>
      <div class="v red">{руб(просроч_всего)} ₽</div></div>
    <div class="card"><div class="k">Доля просрочки</div>
      <div class="v">{просроч_всего / долг_всего * 100:.0f}%</div></div>
    <div class="card"><div class="k">Оплачено за {ОКНО_ОПЛАТ} дн.</div>
      <div class="v">{руб(оплачено_все_окно)} ₽</div>
      <div class="k">{оплат_всего_окно} оплат</div></div>
    <div class="card"><div class="k">Деньги без разноски</div>
      <div class="v amber">{руб(висит_всего)} ₽</div>
      <div class="k">снимет долга {руб(снимет_разноска)} ₽</div></div>
  </div>

  <h2>По ИП</h2>
  <table><thead><tr><th>ИП</th><th class="d">Должников</th><th class="p">Долг, ₽</th>
    <th class="p">Просрочено, ₽</th><th class="d">Доля</th>
    <th class="p">Оплачено за {ОКНО_ОПЛАТ} дн., ₽</th>
    <th class="p">Без разноски, ₽</th></tr></thead>
    <tbody>{сводка_строки}
      <tr class="itogo"><td>ИТОГО</td><td class="d">{len(строки)}</td>
        <td class="p">{руб(долг_всего)}</td><td class="p">{руб(просроч_всего)}</td>
        <td class="d">{просроч_всего / долг_всего * 100:.0f}%</td>
        <td class="p">{руб(оплачено_все_окно)}</td>
        <td class="p">{руб(висит_всего)}</td></tr>
    </tbody></table>
  {двойные_блок}

  <div class="page">
    <h2>Топ просроченных должников ({len(просроченные)} с просрочкой)</h2>
    <div class="note">⚑ — у клиента есть деньги, не разнесённые по накладным:
      сначала разнести, потом звонить. «Платит на карточку …» — деньги приходят на
      другую его карточку (обычно старое юрлицо), а накладные этой не гасятся.</div>
    <table><thead><tr><th class="n">№</th><th>Контрагент</th><th class="ip">ИП</th>
      <th class="p">Просрочено, ₽</th><th class="p">Долг всего, ₽</th>
      <th class="d">Дней</th><th class="pay">Посл. оплата</th></tr></thead>
      <tbody>{таблица_топ(просроченные)}</tbody></table>
  </div>

  <div class="page">
    <h2>Деньги без разноски — долг по ним ложный ({len(без_разноски)})</h2>
    <div class="note">Висит {руб(висит_всего)} ₽; у {должников_с_деньгами} из них есть
      долг — разноска снимет {руб(снимет_разноска)} ₽. Подсвечены те, у кого денег
      хватает на весь долг: им звонить нельзя, нужно разнести оплату.
      «Переплата» — деньги уже разнесены, но на другую накладную (минус по ней отчёт
      не вычитает из долга).</div>
    <table><thead><tr><th>Контрагент</th><th class="ip">ИП</th>
      <th class="p">Денег без разноски, ₽</th><th class="p">В т.ч. переплата, ₽</th>
      <th class="p">Долг, ₽</th><th class="p">Снимет разноска, ₽</th></tr></thead>
      <tbody>{таблица_без_разноски(без_разноски)}</tbody></table>
  </div>

  <div class="page">
    <h2>Последние оплаты — {ОКНО_ПОСЛЕДНИХ} дн. ({len(оплаты_список)} оплат на
      {руб(sum(о["сумма"] for о in оплаты_список))} ₽)</h2>
    <div class="note">Проведённые поступления от покупателей: счёт и касса.
      Подсвечены оплаты в день среза — в долге выше они ещё НЕ учтены.
      {"Показаны первые 150." if len(оплаты_список) > 150 else ""}</div>
    <table><thead><tr><th class="d">Дата</th><th>Контрагент</th><th class="ip">ИП</th>
      <th class="p">Сумма, ₽</th><th class="d">Куда</th>
      <th class="p">Долг сейчас, ₽</th></tr></thead>
      <tbody>{таблица_оплат(оплаты_список, 150)}</tbody></table>
  </div>
  {страницы_ип}
</body></html>
"""

html_path = f"{SCRATCH}/debt.html"
open(html_path, "w", encoding="utf-8").write(HTML)
pdf_path = f"{DESKTOP}/{ИМЯ}.pdf"
chrome = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
subprocess.run([chrome, "--headless", "--disable-gpu", "--no-pdf-header-footer",
                f"--print-to-pdf={pdf_path}", f"file://{html_path}"],
               check=True, capture_output=True, timeout=180)
print("PDF:", pdf_path, os.path.getsize(pdf_path), "байт")
print("должников:", len(строки), "| с просрочкой:", len(просроченные),
      "| долг:", руб(долг_всего), "| просрочено:", руб(просроч_всего))
print("топ-5:", [(s["name"][:30], руб(s["overdue"]), s["days"]) for s in просроченные[:5]])
print(f"оплаты за {ОКНО_ОПЛАТ} дн.: {оплат_всего_окно} шт на {руб(оплачено_все_окно)} ₽"
      f" | за {ОКНО_ПОСЛЕДНИХ} дн.: {len(оплаты_список)} шт")
print(f"деньги без разноски: {руб(висит_всего)} ₽ у {len(без_разноски)} контрагентов,"
      f" из них должников {должников_с_деньгами} — разноска снимет {руб(снимет_разноска)} ₽")
